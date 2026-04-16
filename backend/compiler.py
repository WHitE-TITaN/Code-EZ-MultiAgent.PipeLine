import base64
import hashlib
import io
import os
import random
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:  # pragma: no cover
    Image = None
    ImageDraw = None
    ImageFont = None


SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5

PALETTE = {
    "navy": RGBColor(15, 23, 42),
    "ink": RGBColor(30, 41, 59),
    "slate": RGBColor(71, 85, 105),
    "line": RGBColor(203, 213, 225),
    "white": RGBColor(255, 255, 255),
}

THEMES = {
    "Master1": {"base": RGBColor(252, 228, 228), "glow": RGBColor(255, 204, 204), "corner": RGBColor(253, 215, 215), "accent": RGBColor(229, 35, 50), "accent_dark": RGBColor(15, 15, 15), "band": RGBColor(255, 192, 203), "card_fill": RGBColor(255, 255, 255)},
    "Master2": {"base": RGBColor(248, 250, 252), "glow": RGBColor(254, 226, 226), "corner": RGBColor(241, 245, 249), "accent": RGBColor(239, 76, 76), "accent_dark": RGBColor(30, 41, 59), "band": RGBColor(254, 202, 202), "card_fill": RGBColor(255, 255, 255)},
    "Master3": {"base": RGBColor(235, 242, 237), "glow": RGBColor(218, 232, 222), "corner": RGBColor(204, 222, 209), "accent": RGBColor(76, 175, 80), "accent_dark": RGBColor(27, 59, 36), "band": RGBColor(194, 214, 199), "card_fill": RGBColor(255, 255, 255)},
}

IMAGE_CACHE = {}


def clean_base64_string(b64_string: str) -> str:
    clean = str(b64_string or "").strip()
    if "base64," in clean:
        clean = clean.split("base64,", 1)[1]
    clean = clean.strip().strip("()").strip("'").strip('"').strip()
    clean = re.sub(r"\s+", "", clean)
    clean += "=" * ((4 - len(clean) % 4) % 4)
    return clean


def find_best_match(topic_ref, harvested_keys):
    if not topic_ref:
        return None
    if topic_ref in harvested_keys:
        return topic_ref
    ref = re.sub(r"[^a-z0-9]+", "", str(topic_ref).lower())
    for key in harvested_keys:
        norm = re.sub(r"[^a-z0-9]+", "", str(key).lower())
        if ref == norm or ref in norm or norm in ref:
            return key
    return None


def build_powerpoint(master_blueprint_json: list, harvested_images: dict, output_filename="Final_Deck.pptx", template_name="random"):
    template_choices = sorted(THEMES.keys())
    if not template_name or str(template_name).lower() == "random":
        template_name = random.choice(template_choices)
    template_name = template_name if template_name in THEMES else "Master1"

    template_dir = os.path.join(os.path.dirname(__file__), "Slide Master")
    template_path = os.path.join(template_dir, f"{template_name}.pptx")
    using_template = os.path.exists(template_path)
    prs = Presentation(template_path) if using_template else Presentation()
    if using_template:
        _clear_all_slides(prs)
    if not using_template:
        prs.slide_width = Inches(SLIDE_WIDTH)
        prs.slide_height = Inches(SLIDE_HEIGHT)

    harvested_images = harvested_images or {}
    deck_title = _resolve_deck_title(master_blueprint_json, output_filename)
    if using_template:
        _add_cover_from_template(prs, deck_title, THEMES[template_name])
    else:
        _add_cover_slide(prs, deck_title, using_template=False)

    for index, slide_data in enumerate(master_blueprint_json or [], start=1):
        slide_data = slide_data or {}
        theme = THEMES.get(template_name) if using_template else _resolve_theme(slide_data, index)
        slide = prs.slides.add_slide(_pick_body_layout(prs, using_template))

        _paint_base_background(slide, slide_data, theme, using_template=False)

        image_asset = _resolve_image_asset(slide_data, harvested_images, theme)
        blocks = _normalize_blocks(slide_data)
        layout_type = _normalize_layout_type(slide_data, image_asset)
        _render_layout(slide, layout_type, blocks, image_asset, theme)
        _draw_title(slide, slide_data.get("title", f"Slide {index}"), using_template=False, theme=theme)

    if using_template:
        _add_thank_you_from_template(prs)

    prs.save(output_filename)
    return output_filename


def _resolve_deck_title(master_blueprint_json, output_filename):
    if master_blueprint_json and isinstance(master_blueprint_json[0], dict):
        first_title = str(master_blueprint_json[0].get("title") or "").strip()
        if first_title:
            return first_title
    return os.path.splitext(os.path.basename(output_filename))[0].replace("_", " ").strip() or "Presentation"


def _clear_all_slides(prs):
    slide_ids = list(prs.slides._sldIdLst)
    for slide_id in slide_ids:
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        prs.slides._sldIdLst.remove(slide_id)


def _pick_layout_by_name(prs, keywords, fallback_index=0):
    for index, layout in enumerate(prs.slide_layouts):
        name = (getattr(layout, "name", "") or "").lower()
        if any(keyword in name for keyword in keywords):
            return prs.slide_layouts[index]
    return prs.slide_layouts[min(fallback_index, len(prs.slide_layouts) - 1)]


def _pick_body_layout(prs, using_template):
    if not using_template:
        return prs.slide_layouts[6]
    for keywords in (("blank",), ("title only",), ("title",)):
        layout = _pick_layout_by_name(prs, keywords)
        if layout is not None:
            return layout
    return prs.slide_layouts[0]


def _add_cover_from_template(prs, deck_title, theme):
    slide = prs.slides.add_slide(_pick_layout_by_name(prs, ("cover", "title company"), fallback_index=0))
    _set_native_title_or_textbox(slide, deck_title, top=1.7, height=1.9, font_size=38, is_cover=True, cover_color=theme["accent"])
    return slide


def _add_thank_you_from_template(prs):
    return prs.slides.add_slide(_pick_layout_by_name(prs, ("thank",), fallback_index=max(len(prs.slide_layouts) - 1, 0)))


def _add_cover_slide(prs, deck_title, using_template):
    if using_template:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        _set_native_title(slide, deck_title)
        return slide

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    theme = THEMES["Master2"]
    _paint_base_background(
        slide,
        {"background_design": {"base_fill": "#F4F6FF", "elements": [{"type": "circle", "x": -1.0, "y": -1.0, "w": 4.5, "h": 4.5, "fill": "#CDD6FF", "transparency": 0.18}, {"type": "rounded_rect", "x": 0.9, "y": 1.25, "w": 11.2, "h": 4.85, "fill": "#E9EDFF", "transparency": 0.22}, {"type": "arc_band", "x": 9.3, "y": 4.7, "w": 4.1, "h": 2.8, "fill": "#6E79FF", "transparency": 0.18}]}},
        theme,
        using_template=False,
    )
    box = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(8.2), Inches(2.0))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = deck_title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = PALETTE["navy"]
    return slide


def _set_native_title(slide, title_text):
    if getattr(slide.shapes, "title", None) is not None:
        slide.shapes.title.text = str(title_text or "Untitled Slide")


def _set_native_title_or_textbox(slide, title_text, top=0.72, height=0.95, font_size=24, is_cover=False, cover_color=None):
    placeholder = _find_best_text_placeholder(slide, prefer_cover_placeholder=is_cover)
    if placeholder is not None:
        _apply_text_to_shape(placeholder, title_text, font_size=font_size, is_cover=is_cover, cover_color=cover_color)
        return

    if getattr(slide.shapes, "title", None) is not None:
        _apply_text_to_shape(slide.shapes.title, title_text, font_size=font_size, is_cover=is_cover, cover_color=cover_color)
        return

    box = slide.shapes.add_textbox(Inches(0.85), Inches(top), Inches(11.6), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = str(title_text or "Untitled Slide")
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = cover_color if is_cover and cover_color is not None else PALETTE["navy"]
    p.font.name = "Times New Roman"


def _apply_text_to_shape(shape, text, font_size=24, is_cover=False, cover_color=None):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.text = str(text or "Untitled Slide")
    p.font.size = Pt(font_size)
    p.font.bold = True
    p.font.color.rgb = cover_color if is_cover and cover_color is not None else PALETTE["navy"]
    p.font.name = "Times New Roman"


def _find_best_text_placeholder(slide, prefer_cover_placeholder=False):
    best_shape = None
    best_area = -1
    for shape in slide.shapes:
        if not hasattr(shape, "text_frame"):
            continue
        existing_text = shape.text_frame.text.strip().lower()
        if prefer_cover_placeholder and "click to add text" in existing_text:
            return shape
        if not getattr(shape, "is_placeholder", False):
            continue
        area = int(shape.width) * int(shape.height)
        if area > best_area:
            best_area = area
            best_shape = shape
    return best_shape


def _resolve_theme(slide_data, slide_index):
    explicit = str(slide_data.get("theme_variant") or "").strip().lower()
    if explicit in THEMES:
        return THEMES[explicit]
    names = sorted(THEMES.keys())
    seed = int(hashlib.sha256(f"{slide_index}:{slide_data.get('title','')}".encode("utf-8")).hexdigest()[:8], 16)
    return THEMES[names[seed % len(names)]]


def _resolve_image_asset(slide_data, harvested_images, theme):
    existing_chart = slide_data.get("existing_chart", {}) or {}
    topic_ref = existing_chart.get("chart_topic_reference")
    best_match_key = find_best_match(topic_ref, list(harvested_images.keys()))
    if existing_chart.get("has_existing_chart") and best_match_key:
        matched_asset = harvested_images.get(best_match_key)
        table_asset = _extract_table_asset(matched_asset)
        if table_asset:
            return {"stream": None, "kind": "table", "table": table_asset}
        stream = _first_image_stream(matched_asset)
        if stream:
            return {"stream": stream, "kind": "chart"}

    foreground = slide_data.get("foreground_image", {}) or {}
    prompt = foreground.get("foreground_image_prompt") or _build_visual_prompt(slide_data)
    if foreground.get("has_foreground_image") or str(slide_data.get("visual_priority") or "").lower() in {"image_focus", "balanced"}:
        stream = _generate_local_visual(prompt, 1024, 768, theme)
        if stream:
            return {"stream": stream, "kind": "illustration"}
    return {"stream": None, "kind": "none"}


def _extract_table_asset(raw_value):
    matrix = _coerce_table_matrix(raw_value)
    if matrix:
        return matrix
    if isinstance(raw_value, dict):
        for key in ("table", "table_data", "rows", "data", "content", "text", "markdown"):
            if raw_value.get(key):
                matrix = _coerce_table_matrix(raw_value.get(key))
                if matrix:
                    return matrix
    if isinstance(raw_value, (list, tuple)):
        for item in raw_value:
            matrix = _coerce_table_matrix(item)
            if matrix:
                return matrix
    return None


def _coerce_table_matrix(value):
    if value is None:
        return None
    if isinstance(value, dict):
        headers = value.get("headers") or value.get("columns")
        rows = value.get("rows") or value.get("values") or value.get("data")
        if isinstance(headers, (list, tuple)) and isinstance(rows, (list, tuple)):
            matrix = [[str(cell).strip() for cell in headers]]
            for row in rows:
                if isinstance(row, dict):
                    matrix.append([str(row.get(str(header), row.get(header, ""))).strip() for header in headers])
                elif isinstance(row, (list, tuple)):
                    matrix.append([str(cell).strip() for cell in row])
            return _clean_table_matrix(matrix)
        for key in ("table", "table_data", "content", "text", "markdown"):
            if value.get(key):
                matrix = _coerce_table_matrix(value.get(key))
                if matrix:
                    return matrix
        return None
    if isinstance(value, str):
        return _parse_markdown_table(value)
    if isinstance(value, (list, tuple)):
        if value and all(isinstance(item, dict) for item in value):
            headers = []
            for item in value:
                for key in item.keys():
                    if key not in headers:
                        headers.append(str(key))
            matrix = [headers]
            for item in value:
                matrix.append([str(item.get(header, "")).strip() for header in headers])
            return _clean_table_matrix(matrix)
        if value and all(isinstance(item, (list, tuple)) for item in value):
            matrix = [[str(cell).strip() for cell in row] for row in value]
            return _clean_table_matrix(matrix)
    return None


def _parse_markdown_table(text):
    lines = [line.strip() for line in str(text or "").splitlines() if "|" in line]
    if len(lines) < 2:
        return None
    matrix = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if cells:
            matrix.append(cells)
    if len(matrix) < 2:
        return None
    separator = matrix[1]
    if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in separator if cell):
        matrix.pop(1)
    return _clean_table_matrix(matrix)


def _clean_table_matrix(matrix):
    if not isinstance(matrix, list) or len(matrix) < 2:
        return None
    normalized = []
    max_cols = 0
    for row in matrix:
        if not isinstance(row, list):
            return None
        cleaned = [str(cell or "").strip() for cell in row]
        if any(cleaned):
            normalized.append(cleaned)
            max_cols = max(max_cols, len(cleaned))
    if len(normalized) < 2 or max_cols < 2:
        return None
    return [row + [""] * (max_cols - len(row)) for row in normalized[:8]]


def _first_image_stream(raw_value):
    if raw_value is None:
        return None
    if isinstance(raw_value, (list, tuple)):
        for item in raw_value:
            stream = _coerce_image_stream(item)
            if stream:
                return stream
        return None
    return _coerce_image_stream(raw_value)


def _coerce_image_stream(image_value):
    if image_value is None:
        return None
    if isinstance(image_value, io.BytesIO):
        image_value.seek(0)
        return io.BytesIO(image_value.read())
    if isinstance(image_value, bytes):
        return io.BytesIO(image_value)
    if isinstance(image_value, dict):
        for key in ("data", "base64", "image_base64", "content", "path", "url"):
            if image_value.get(key):
                stream = _coerce_image_stream(image_value[key])
                if stream:
                    return stream
        return None
    if not isinstance(image_value, str):
        return None
    candidate = image_value.strip()
    if candidate.startswith("(") and candidate.endswith(")"):
        candidate = candidate[1:-1].strip()
    if candidate.startswith("data:image/") or "base64," in candidate:
        try:
            return io.BytesIO(base64.b64decode(clean_base64_string(candidate)))
        except Exception:
            return None
    if os.path.exists(candidate):
        try:
            with open(candidate, "rb") as f:
                return io.BytesIO(f.read())
        except OSError:
            return None
    try:
        return io.BytesIO(base64.b64decode(clean_base64_string(candidate)))
    except Exception:
        return None


def _normalize_blocks(slide_data):
    blocks = slide_data.get("content_blocks")
    if isinstance(blocks, list) and blocks:
        out = []
        for block in blocks:
            if isinstance(block, dict):
                bullets = block.get("bullets", [])
                if isinstance(bullets, str):
                    bullets = [bullets]
                out.append({"block_title": str(block.get("block_title") or "Insight"), "body": str(block.get("body") or ""), "bullets": [str(x).strip() for x in bullets if str(x).strip()]})
        if out:
            return out

    summary = str(slide_data.get("summary") or "")
    bullets = [str(x).strip() for x in (slide_data.get("bullets") or []) if str(x).strip()]
    metrics = [str(x).strip() for x in (slide_data.get("highlight_metrics") or []) if str(x).strip()]
    out = []
    if summary or bullets:
        out.append({"block_title": "Key Takeaway", "body": summary, "bullets": bullets[:4]})
    if metrics:
        out.append({"block_title": "Key Metrics", "body": "", "bullets": metrics[:4]})
    return out or [{"block_title": "Overview", "body": "", "bullets": []}]


def _normalize_layout_type(slide_data, image_asset):
    layout_type = str(slide_data.get("layout_type") or "split_left_image_right_cards")
    if image_asset.get("kind") in {"chart", "table"}:
        return "split_left_image_right_cards"
    if layout_type not in {"top_image_bottom_cards", "hero_background_with_floating_cards", "split_left_image_right_cards", "diagonal_text_split", "multi_column_cards"}:
        return "split_left_image_right_cards"
    return layout_type


def _render_layout(slide, layout_type, blocks, image_asset, theme):
    if layout_type == "top_image_bottom_cards":
        _layout_top_image_bottom_cards(slide, blocks, image_asset, theme)
    elif layout_type == "hero_background_with_floating_cards":
        _layout_hero_background_with_floating_cards(slide, blocks, image_asset, theme)
    elif layout_type == "multi_column_cards":
        _layout_multi_column_cards(slide, blocks, image_asset, theme)
    elif layout_type == "diagonal_text_split":
        _layout_diagonal_text_split(slide, blocks, image_asset, theme)
    else:
        _layout_split_left_image_right_cards(slide, blocks, image_asset, theme)


def _layout_top_image_bottom_cards(slide, blocks, image_asset, theme):
    left, top, width, height = _content_bounds()
    hero_height = 2.45
    card_top = top + hero_height + 0.26
    card_height = 2.34
    if image_asset.get("kind") == "table" and image_asset.get("table"):
        _add_table_visual(slide, image_asset["table"], Inches(left), Inches(top), Inches(width), Inches(hero_height), theme)
    elif image_asset.get("stream"):
        _add_picture_cover(slide, image_asset["stream"], Inches(left), Inches(top), Inches(width), Inches(hero_height))
    else:
        _add_decorative_visual_panel(slide, theme, left=Inches(left + 4.0), top=Inches(top + 0.15), width=Inches(3.2), height=Inches(1.95))
    for index, (card_left, card_top_in, card_width, card_height_in) in enumerate(_horizontal_specs(len(blocks), Inches(card_top), Inches(card_height), gutter=0.24, left=left, usable_width=width)):
        _add_card(slide, card_left, card_top_in, card_width, card_height_in, blocks[index], theme)


def _layout_hero_background_with_floating_cards(slide, blocks, image_asset, theme):
    left, top, width, height = _content_bounds()
    visual_width = 4.55
    visual_height = 2.55
    visual_left = left + width - visual_width
    if image_asset.get("kind") == "table" and image_asset.get("table"):
        _add_table_visual(slide, image_asset["table"], Inches(visual_left), Inches(top), Inches(visual_width), Inches(visual_height), theme)
    elif image_asset.get("stream"):
        _add_picture_cover(slide, image_asset["stream"], Inches(visual_left), Inches(top), Inches(visual_width), Inches(visual_height))
    else:
        _add_decorative_visual_panel(slide, theme, left=Inches(visual_left + 0.35), top=Inches(top + 0.15), width=Inches(3.5), height=Inches(2.0))
    specs = [
        (Inches(left), Inches(top + 0.05), Inches(4.7), Inches(2.12)),
        (Inches(left + 0.25), Inches(top + 2.48), Inches(4.15), Inches(1.92)),
        (Inches(left + 4.72), Inches(top + 2.68), Inches(3.55), Inches(1.78)),
        (Inches(visual_left), Inches(top + 2.48), Inches(4.55), Inches(1.92)),
    ]
    for index, block in enumerate(blocks[: len(specs)]):
        card_left, card_top, card_width, card_height = specs[index]
        _add_card(slide, card_left, card_top, card_width, card_height, block, theme, dark=True)


def _layout_multi_column_cards(slide, blocks, image_asset, theme):
    left, top, width, height = _content_bounds()
    _add_band(slide, Inches(left), Inches(top), Inches(width), Inches(height), theme["band"], 0.18)
    visual_width = 4.8
    column_gap = 0.26
    cards_left = left + 0.14
    cards_top = top + 0.14
    cards_width = width - visual_width - column_gap - 0.24
    cards_height = height - 0.28
    visual_left = left + width - visual_width - 0.14
    visual_top = top + 0.14
    visual_height = height - 0.28
    card_specs = _stacked_specs(cards_left, cards_top, cards_width, cards_height, min(len(blocks), 3), gutter=0.18)
    if image_asset.get("kind") == "table" and image_asset.get("table"):
        _add_visual_frame(slide, Inches(visual_left), Inches(visual_top), Inches(visual_width), Inches(visual_height), theme)
        _add_table_visual(slide, image_asset["table"], Inches(visual_left + 0.12), Inches(visual_top + 0.14), Inches(visual_width - 0.24), Inches(visual_height - 0.28), theme)
        for index, block in enumerate(blocks[: len(card_specs)]):
            card_left, card_top, card_width, card_height = card_specs[index]
            _add_card(slide, card_left, card_top, card_width, card_height, block, theme)
        return
    if image_asset.get("stream") and image_asset.get("kind") != "chart":
        _add_visual_frame(slide, Inches(visual_left), Inches(visual_top), Inches(visual_width), Inches(visual_height), theme)
        _add_picture_cover(slide, image_asset["stream"], Inches(visual_left), Inches(visual_top), Inches(visual_width), Inches(visual_height))
        for index, block in enumerate(blocks[: len(card_specs)]):
            card_left, card_top, card_width, card_height = card_specs[index]
            _add_card(slide, card_left, card_top, card_width, card_height, block, theme)
        return
    _add_decorative_visual_panel(slide, theme, left=Inches(visual_left + 0.55), top=Inches(visual_top + 0.5), width=Inches(3.45), height=Inches(2.15))
    for index, (card_left, card_top, card_width, card_height) in enumerate(_horizontal_specs(len(blocks), Inches(top + 0.34), Inches(height - 0.68), gutter=0.24, left=left + 0.24, usable_width=width - 0.48)):
        _add_card(slide, card_left, card_top, card_width, card_height, blocks[index], theme)


def _layout_diagonal_text_split(slide, blocks, image_asset, theme):
    left, top, width, height = _content_bounds()
    if blocks:
        _add_card(slide, Inches(left), Inches(top), Inches(4.35), Inches(2.15), blocks[0], theme)
    if len(blocks) > 1:
        _add_card(slide, Inches(left + width - 4.35), Inches(top + height - 2.15), Inches(4.35), Inches(2.15), blocks[1], theme)
    if len(blocks) > 2:
        _add_card(slide, Inches(left + 0.2), Inches(top + height - 1.72), Inches(3.9), Inches(1.5), blocks[2], theme, dark=True)
    if image_asset.get("kind") == "table" and image_asset.get("table"):
        _add_table_visual(slide, image_asset["table"], Inches(left + 4.45), Inches(top + 0.55), Inches(4.3), Inches(3.15), theme)
    elif image_asset.get("stream"):
        _add_picture_contain(slide, image_asset["stream"], Inches(left + 4.45), Inches(top + 0.55), Inches(4.3), Inches(3.15))
    else:
        _add_decorative_visual_panel(slide, theme, left=Inches(left + 4.95), top=Inches(top + 0.75), width=Inches(3.45), height=Inches(2.15))


def _layout_split_left_image_right_cards(slide, blocks, image_asset, theme):
    left, top, width, height = _content_bounds()
    split_gap = 0.28
    visual_width = 6.15
    visual_height = height
    visual_left = left
    cards_left = visual_left + visual_width + split_gap
    cards_width = width - visual_width - split_gap
    card_specs = _stacked_specs(cards_left, top, cards_width, height, min(len(blocks), 3), gutter=0.16)
    if image_asset.get("kind") == "table" and image_asset.get("table"):
        _add_visual_frame(slide, Inches(visual_left), Inches(top), Inches(visual_width), Inches(visual_height), theme)
        _add_table_visual(slide, image_asset["table"], Inches(visual_left + 0.12), Inches(top + 0.16), Inches(visual_width - 0.24), Inches(visual_height - 0.32), theme)
    elif image_asset.get("stream"):
        _add_visual_frame(slide, Inches(visual_left), Inches(top), Inches(visual_width), Inches(visual_height), theme)
        _add_picture_cover(slide, image_asset["stream"], Inches(visual_left), Inches(top), Inches(visual_width), Inches(visual_height))
    else:
        _add_decorative_visual_panel(slide, theme, left=Inches(visual_left + 0.9), top=Inches(top + 0.8), width=Inches(4.8), height=Inches(2.65))
    for index, (card_left, card_top, card_width, card_height) in enumerate(card_specs):
        _add_card(slide, Inches(card_left), Inches(card_top), Inches(card_width), Inches(card_height), blocks[index], theme)


def _content_bounds():
    return (0.62, 1.42, 12.08, 4.98)


def _horizontal_specs(count, top, height, gutter=0.28, left=0.62, usable_width=12.05):
    count = max(1, min(count, 4))
    usable = usable_width - (gutter * (count - 1))
    width = usable / count
    current_left = left
    specs = []
    for _ in range(count):
        specs.append((Inches(current_left), top, Inches(width), height))
        current_left += width + gutter
    return specs


def _stacked_specs(left, top, width, height, count, gutter=0.18):
    count = max(1, count)
    usable_height = height - (gutter * (count - 1))
    card_height = usable_height / count
    specs = []
    current_top = top
    for _ in range(count):
        specs.append((left, current_top, width, card_height))
        current_top += card_height + gutter
    return specs


def _paint_base_background(slide, slide_data, theme, using_template=False):
    if using_template:
        return
    base_fill = _coerce_rgb((slide_data.get("background_design") or {}).get("base_fill")) or theme["base"]
    base = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(SLIDE_WIDTH), Inches(SLIDE_HEIGHT))
    base.fill.solid()
    base.fill.fore_color.rgb = base_fill
    base.line.fill.background()
    elements = (slide_data.get("background_design") or {}).get("elements")
    if isinstance(elements, list) and elements:
        _render_background_design(slide, elements, theme)
    else:
        glow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(-0.9), Inches(-1.1), Inches(4.2), Inches(4.2))
        glow.fill.solid()
        glow.fill.fore_color.rgb = theme["glow"]
        glow.fill.transparency = 0.45
        glow.line.fill.background()


def _render_background_design(slide, elements, theme):
    for element in elements[:6]:
        if not isinstance(element, dict):
            continue
        x = Inches(float(element.get("x", 0)))
        y = Inches(float(element.get("y", 0)))
        w = Inches(max(0.1, float(element.get("w", 1.0))))
        h = Inches(max(0.1, float(element.get("h", 1.0))))
        fill = _coerce_rgb(element.get("fill")) or theme["glow"]
        transparency = max(0.0, min(0.95, float(element.get("transparency", 0.2))))
        kind = str(element.get("type") or "").lower()
        if kind == "circle":
            shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, x, y, w, h)
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
            shape.fill.transparency = transparency
            shape.line.fill.background()
        elif kind == "arc_band":
            shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ARC, x, y, w, h)
            shape.fill.background()
            shape.line.color.rgb = fill
            shape.line.width = Pt(10)
        else:
            auto_shape = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if kind == "rounded_rect" else MSO_AUTO_SHAPE_TYPE.RECTANGLE
            shape = slide.shapes.add_shape(auto_shape, x, y, w, h)
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
            shape.fill.transparency = transparency
            shape.line.fill.background()


def _coerce_rgb(value):
    if isinstance(value, RGBColor):
        return value
    text = str(value or "").strip().lstrip("#")
    if len(text) != 6:
        return None
    try:
        return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))
    except ValueError:
        return None


def _draw_title(slide, title_text, using_template, theme=None):
    if using_template:
        return
    theme = theme or THEMES["Master1"]

    title_back = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(0.58),
        Inches(0.46),
        Inches(11.75),
        Inches(0.92),
    )
    title_back.fill.solid()
    title_back.fill.fore_color.rgb = PALETTE["white"]
    title_back.fill.transparency = 0.18
    title_back.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.78), Inches(0.62), Inches(11.7), Inches(0.95))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(title_text or "Untitled Slide")
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = theme["accent_dark"]

    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.62),
        Inches(0.64),
        Inches(1.9),
        Inches(0.08),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = theme["accent"]
    accent.line.fill.background()


def _add_card(slide, left, top, width, height, block_data, theme, dark=False):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = theme["card_fill"]
    card.fill.transparency = 0.02
    card.line.color.rgb = theme["accent"] if dark else PALETTE["line"]
    card.line.width = Pt(1.15)
    _draw_text_block(card, block_data, dark)
    return card


def _draw_text_block(shape, block_data, dark=False):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.14)
    tf.margin_bottom = Inches(0.12)
    title_para = tf.paragraphs[0]
    title_para.text = str(block_data.get("block_title") or "Insight")
    title_para.font.size = Pt(15)
    title_para.font.bold = True
    title_para.font.color.rgb = PALETTE["navy"]
    title_para.font.name = "Times New Roman"
    title_para.space_after = Pt(6)
    body = str(block_data.get("body") or "").strip()
    if body:
        p = tf.add_paragraph()
        p.text = body
        p.font.size = Pt(14)
        p.font.color.rgb = PALETTE["slate"]
        p.font.name = "Times New Roman"
        p.space_after = Pt(5)
    for bullet in (block_data.get("bullets") or [])[:6]:
        p = tf.add_paragraph()
        p.text = f"- {str(bullet)}"
        p.font.size = Pt(14)
        p.font.color.rgb = PALETTE["ink"]
        p.font.name = "Times New Roman"
        p.space_after = Pt(3)


def _add_band(slide, left, top, width, height, color, transparency):
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    band.fill.solid()
    band.fill.fore_color.rgb = color
    band.fill.transparency = transparency
    band.line.fill.background()
    return band


def _add_visual_frame(slide, left, top, width, height, theme):
    shadow = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left + Inches(0.08), top + Inches(0.1), width, height)
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = RGBColor(148, 163, 184)
    shadow.fill.transparency = 0.72
    shadow.line.fill.background()
    frame = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    frame.fill.solid()
    frame.fill.fore_color.rgb = PALETTE["white"]
    frame.line.color.rgb = theme["glow"]
    frame.line.width = Pt(1.2)
    return frame


def _add_decorative_visual_panel(slide, theme, left=Inches(9.0), top=Inches(1.55), width=Inches(2.75), height=Inches(1.95)):
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = theme["glow"]
    panel.fill.transparency = 0.18
    panel.line.fill.background()
    orbit = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, left + Inches(0.45), top + Inches(0.3), Inches(1.1), Inches(1.1))
    orbit.fill.solid()
    orbit.fill.fore_color.rgb = theme["accent"]
    orbit.fill.transparency = 0.08
    orbit.line.fill.background()


def _add_picture_cover(slide, image_stream, left, top, width, height):
    return _add_picture(slide, image_stream, left, top, width, height, "cover")


def _add_picture_contain(slide, image_stream, left, top, width, height):
    return _add_picture(slide, image_stream, left, top, width, height, "contain")


def _add_table_visual(slide, matrix, left, top, width, height, theme):
    if not matrix or len(matrix) < 2:
        return None
    rows = len(matrix)
    cols = max(len(row) for row in matrix)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    table.first_row = True
    column_width = int(width / max(cols, 1))
    row_height = int(height / max(rows, 1))
    for col_index in range(cols):
        table.columns[col_index].width = column_width
    for row_index in range(rows):
        table.rows[row_index].height = row_height
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            cell.text = matrix[row_index][col_index] if col_index < len(matrix[row_index]) else ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme["accent"] if row_index == 0 else theme["card_fill"]
            cell.text_frame.word_wrap = True
            cell.text_frame.margin_left = Inches(0.08)
            cell.text_frame.margin_right = Inches(0.08)
            cell.text_frame.margin_top = Inches(0.04)
            cell.text_frame.margin_bottom = Inches(0.04)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER if row_index == 0 else PP_ALIGN.LEFT
                if paragraph.runs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(12.5 if row_index == 0 else 11.5)
                        run.font.bold = row_index == 0
                        run.font.color.rgb = PALETTE["white"] if row_index == 0 else theme["accent_dark"]
            for border_name in ("border_left", "border_right", "border_top", "border_bottom"):
                border = getattr(cell, border_name, None)
                if border is not None and hasattr(border, "color"):
                    border.color.rgb = theme["band"]
    return table_shape


def _add_picture(slide, image_stream, left, top, width, height, mode):
    image_bytes = _read_stream_bytes(image_stream)
    if not image_bytes:
        return None
    image_size = _read_image_size(image_bytes)
    if not image_size:
        return slide.shapes.add_picture(io.BytesIO(image_bytes), left, top, width=width, height=height)
    image_ratio = image_size[0] / image_size[1]
    frame_ratio = width / height
    if mode == "cover":
        draw_width, draw_height = (height * image_ratio, height) if image_ratio > frame_ratio else (width, width / image_ratio)
    else:
        draw_width, draw_height = (width, width / image_ratio) if image_ratio > frame_ratio else (height * image_ratio, height)
    draw_left = left + (width - draw_width) / 2
    draw_top = top + (height - draw_height) / 2
    return slide.shapes.add_picture(io.BytesIO(image_bytes), draw_left, draw_top, width=draw_width, height=draw_height)


def _read_stream_bytes(stream):
    if isinstance(stream, io.BytesIO):
        stream.seek(0)
        return stream.read()
    if isinstance(stream, bytes):
        return stream
    return None


def _read_image_size(image_bytes):
    if Image is None:
        return None
    try:
        with Image.open(io.BytesIO(image_bytes)) as image:
            return image.size
    except Exception:
        return None


def _generate_local_visual(prompt, width, height, theme):
    cache_key = (str(prompt), width, height, str(theme["accent"]))
    if cache_key in IMAGE_CACHE:
        return io.BytesIO(IMAGE_CACHE[cache_key])
    if Image is None or ImageDraw is None:
        return None
    image_width = max(800, int(width))
    image_height = max(450, int(height))
    rng = random.Random(int(hashlib.sha256(str(prompt).encode("utf-8")).hexdigest()[:16], 16))
    base_rgb = _rgb_tuple(theme["base"])
    glow_rgb = _rgb_tuple(theme["glow"])
    accent_rgb = _rgb_tuple(theme["accent"])
    corner_rgb = _rgb_tuple(theme["corner"])
    dark_rgb = _rgb_tuple(theme["accent_dark"])

    canvas = Image.new("RGB", (image_width, image_height), base_rgb)
    draw = ImageDraw.Draw(canvas, "RGBA")

    for y in range(image_height):
        mix = y / max(image_height - 1, 1)
        start = glow_rgb if mix < 0.5 else corner_rgb
        end = corner_rgb if mix < 0.5 else base_rgb
        local_mix = mix * 2 if mix < 0.5 else (mix - 0.5) * 2
        r = int(start[0] + (end[0] - start[0]) * local_mix)
        g = int(start[1] + (end[1] - start[1]) * local_mix)
        b = int(start[2] + (end[2] - start[2]) * local_mix)
        draw.line((0, y, image_width, y), fill=(r, g, b, 255))
    for _ in range(10):
        x0 = rng.randint(-image_width // 8, image_width)
        y0 = rng.randint(-image_height // 8, image_height)
        w = rng.randint(image_width // 8, image_width // 2)
        h = rng.randint(image_height // 8, image_height // 2)
        fill = rng.choice(
            [
                (glow_rgb[0], glow_rgb[1], glow_rgb[2], 55),
                (accent_rgb[0], accent_rgb[1], accent_rgb[2], 42),
                (corner_rgb[0], corner_rgb[1], corner_rgb[2], 48),
                (base_rgb[0], base_rgb[1], base_rgb[2], 42),
            ]
        )
        draw.rounded_rectangle((x0, y0, x0 + w, y0 + h), radius=rng.randint(18, 60), fill=fill)

    draw.rounded_rectangle(
        (40, 40, image_width - 40, image_height - 40),
        radius=28,
        outline=(accent_rgb[0], accent_rgb[1], accent_rgb[2], 90),
        width=3,
    )
    stream = io.BytesIO()
    canvas.save(stream, format="PNG")
    stream.seek(0)
    IMAGE_CACHE[cache_key] = stream.getvalue()
    return io.BytesIO(stream.getvalue())


def _build_visual_prompt(slide_data):
    title = str(slide_data.get("title") or "Business strategy").strip()
    blocks = slide_data.get("content_blocks") or []
    parts = []
    for block in blocks[:2]:
        if isinstance(block, dict):
            parts.append(str(block.get("block_title") or "").strip())
    joined = ", ".join([item for item in parts if item])
    return f"{title}, {joined}" if joined else title


def _rgb_tuple(color):
    return (color[0], color[1], color[2])
